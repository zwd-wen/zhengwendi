import os
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SHOTS = r"c:\Users\lenovo\.trae-cn\builtin\work\shots"
OUT_DOC = r"d:\高级网页设计\1\525-郑雯迪.docx"

# ---------- 生成测试截图 ----------
desk = Image.open(os.path.join(SHOTS, "ch_desktop_full.png"))
edge = Image.open(os.path.join(SHOTS, "edge_full.png"))
mob = Image.open(os.path.join(SHOTS, "ch_mobile_full.png"))
desk.crop((0, 0, 1200, 900)).save(os.path.join(SHOTS, "test_desktop_chrome.png"))
edge.crop((0, 0, 1200, 900)).save(os.path.join(SHOTS, "test_desktop_edge.png"))
mob.crop((0, 0, 390, 844)).save(os.path.join(SHOTS, "test_mobile_chrome.png"))

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.4)
sec.right_margin = Cm(2.4)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)


def run_font(run, cn_font="宋体", size=10.5, bold=False, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(text, size=10.5, bold=False, cn_font="宋体", align=None, first_indent=True,
         space_before=0, space_after=0, line_spacing=1.5, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if first_indent:
        pf.first_line_indent = Pt(size * 2)
    r = p.add_run(text)
    run_font(r, cn_font, size, bold, color)
    return p


def heading(text, size=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    run_font(r, "黑体", size, True)
    return p


def sub(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    r = p.add_run(text)
    run_font(r, "黑体", 11, True)
    return p


def shade(p, fill="F4F1EA"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def code_block(code):
    for line in code.split("\n"):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.05
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.3)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(7.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        shade(p)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 1.0


def picture(path, width_cm, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    c.paragraph_format.line_spacing = 1.0
    r = c.add_run(caption)
    run_font(r, "宋体", 9, False, (110, 110, 110))


def label(text, before=8):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(4)
    pf.keep_with_next = True
    r = p.add_run(text)
    run_font(r, "宋体", 10.5, True)
    return p


# ---------- 封面信息 ----------
t1 = doc.add_paragraph()
t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
t1.paragraph_format.space_before = Pt(10)
t1.paragraph_format.space_after = Pt(2)
r = t1.add_run("《高级网页设计I》")
run_font(r, "黑体", 20, True)

t2 = doc.add_paragraph()
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
t2.paragraph_format.space_after = Pt(16)
r = t2.add_run("网页元素设计 项目报告")
run_font(r, "黑体", 22, True)

t3 = doc.add_paragraph()
t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
t3.paragraph_format.space_after = Pt(14)
r = t3.add_run("学  号：525          姓  名：郑雯迪")
run_font(r, "宋体", 12)

table = doc.add_table(rows=2, cols=2)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(4.2), Cm(11.8)]
rows_data = [("班级", "计算机24005"), ("项目选题名称", "四级项目一 · 网页元素设计（校园读书节专题页）")]
for i, (k, v) in enumerate(rows_data):
    for j, txt in enumerate((k, v)):
        cell = table.cell(i, j)
        cell.width = widths[j]
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
        cp.paragraph_format.space_before = Pt(3)
        cp.paragraph_format.space_after = Pt(3)
        cp.paragraph_format.line_spacing = 1.2
        cr = cp.add_run(txt)
        run_font(cr, "宋体", 10.5, j == 0)

# ---------- 一、实施过程 ----------
heading("实施过程")

sub("项目需求与实现")
para("本项目依据实验指导书要求，从三级项目的选题中选取三个典型网页元素——上图下文、左图右文、列表，使用 HTML 与 CSS 完成静态显示效果设计，不涉及交互与动画。项目选题为“校园读书节专题页”，页面文件为 四级项目1.html，三个元素依次为“本期书单（上图下文）”“本周共读（左图右文）”“活动日程（列表）”，其中前两个元素使用真实图片，左图右文还包含一个纯 CSS 绘制的条形图（图表元素）。")
para("实现前先完成测量：以 1200px 视口为基准，内容区宽度 1124px，卡片三列布局，单卡宽 360px，列间距 22px，卡片配图统一按 3:2 裁切；左图右文两栏宽度比约为 1.02 : 1，栏间距 50px，配图按 4:3 裁切；列表行高约 100px，日期栏宽 84px。配色沿用作品集全局色板：米色底 #f5f1e8、卡片底 #fffdf8、珊瑚红 #d95d46、深蓝 #245b72、暖黄 #f2c05c。")
para("样式设计上先写全局 CSS：用 :root 定义颜色、字体等自定义属性，重置 margin、padding 与盒模型，统一标题、图片、链接的基础样式；再分别为三个元素编写布局样式。布局只使用 CSS Grid 与 Flex，不使用 float，也不使用表格布局。所有类名、自定义属性均以 yl_ 开头，样式中不使用 id 选择器，id 仅用于区块的语义标记与锚点。图标统一调用 font 目录下的 iconfont 字体图标，页面不引入任何第三方库与 CDN。")

label("网页元素1原图（自绘设计稿）：")
picture(os.path.join(SHOTS, "orig1.png"), 14.4, "图 1-1  上图下文 · 设计稿（草测宽度 1122 × 745）")

label("网页元素1实现图：", 4)
picture(os.path.join(SHOTS, "impl_yl_sec_cards.png"), 14.4, "图 1-2  上图下文 · 实现效果（Chrome，1200 视口）")
para("对比说明：三张卡片的宽度、间距、图片比例与设计稿一致，图片在上、文字在下；标题、作者、简介与收藏信息自上而下排列，左右两列卡片底边对齐。设计稿中的灰色占位块在实现时替换为拍摄素材，图标位置改用 iconfont 字体的“喜欢”“收藏”图标。")

label("网页元素2原图（自绘设计稿）：")
picture(os.path.join(SHOTS, "orig2.png"), 14.4, "图 2-1  左图右文 · 设计稿（草测宽度 1122 × 669）")

label("网页元素2实现图：", 4)
picture(os.path.join(SHOTS, "impl_yl_sec_feature.png"), 14.4, "图 2-2  左图右文 · 实现效果（Chrome，1200 视口）")
para("对比说明：左侧配图、右侧文字的 1.02 : 1 两栏结构与设计稿对应，右栏依次为栏目眉标、两行主标题、说明文字、三项数据、进度条形图，垂直方向与左图居中对齐。条形图用五个 span 的百分比宽度绘制，并用 nth-child 为不同章节区分红、蓝、黄三色，与设计稿中的灰色占位条相比信息更完整。")

label("网页元素3原图（自绘设计稿）：")
picture(os.path.join(SHOTS, "orig3.png"), 14.4, "图 3-1  列表 · 设计稿（草测宽度 1122 × 700）")

label("网页元素3实现图：", 4)
picture(os.path.join(SHOTS, "impl_yl_sec_list.png"), 14.4, "图 3-2  列表 · 实现效果（Chrome，1200 视口）")
para("对比说明：列表采用有序列表 ol 承载四场活动，每行按“日期 84px + 正文自适应 + 状态标签”三列排列，行间以 1px 细线分隔，与设计稿的行结构与间距一致。日期中的日号使用衬线字体放大，状态标签分为“可报名”（实心珊瑚红）与“已约满 / 筹备中”（描边）两种，图标取自 iconfont。")

# ---------- 二、项目代码 ----------
heading("项目代码")
para("代码遵循“全局 CSS + class 选择器”的组织方式：:root 中的全局变量统一定义配色与字体，页面所有样式类名均以 yl_ 开头，未使用 id 选择器，也未使用表格布局。页面 head 注释中写有姓名与日期，title 与页脚中写有学号 525 和姓名，满足代码中包含学号或姓名的要求。以下为三个元素的核心代码。")

sub("网页元素1核心代码（上图下文）")
code_block('''<!-- HTML 结构 -->
<div class="yl_cards">
    <article class="yl_card">
        <figure class="yl_card_media">
            <img src="img/yl_book1.jpg" alt="窗台上摆放的一摞硬皮书与陶制茶杯">
        </figure>
        <div class="yl_card_body">
            <span class="yl_tag">城市 · 随笔</span>
            <h3>城市与记忆</h3>
            <p class="yl_author">林知远 / 著 · 2024</p>
            <p class="yl_card_text">沿着旧城区的街道行走，把一座城市拆成气味、声音与光线……</p>
            <div class="yl_meta">
                <span><i class="iconfont icon-xihuan"></i>1,286</span>
                <span><i class="iconfont icon-shoucang"></i>已收藏</span>
            </div>
        </div>
    </article>
    <!-- 另两张卡片结构相同 -->
</div>

/* CSS：三列栅格 + 上图下文 */
.yl_cards { display: grid; grid-template-columns: repeat(3, 1fr); gap: 22px; }
.yl_card {
    display: flex; flex-direction: column; overflow: hidden;
    background: var(--yl_card); border: 1px solid var(--yl_line);
}
.yl_card_media { aspect-ratio: 3 / 2; overflow: hidden; }
.yl_card_media img { width: 100%; height: 100%; object-fit: cover; }
.yl_card_body { display: flex; flex: 1; flex-direction: column; padding: 22px 22px 24px; }
.yl_meta { display: flex; gap: 18px; margin-top: auto; padding-top: 18px;
    border-top: 1px dashed var(--yl_line); color: var(--yl_blue); font-size: 12px; }''')

sub("网页元素2核心代码（左图右文 + CSS 图表）")
code_block('''<!-- HTML 结构：左图右文 -->
<section class="yl_section" id="yl_sec_feature">
    <div class="yl_feature">
        <figure class="yl_feature_media">
            <img src="img/yl_book4.jpg" alt="临海窗边圆桌上的书、茶杯与一摞书">
            <figcaption>FIG. 02 — 本周共读现场 · 海边的阅读角</figcaption>
        </figure>
        <div class="yl_feature_body">
            <span class="yl_no">02 / ELEMENT</span>
            <h2>左图右文<br>把一本书读成一段旅程</h2>
            <div class="yl_stats">…三项数据…</div>
            <div class="yl_chart" role="img" aria-label="各章节阅读完成度条形图">
                <div class="yl_chart_row">
                    <span class="yl_chart_label">第一章 出发</span>
                    <span class="yl_chart_track">
                        <span class="yl_chart_fill" style="width: 92%"></span>
                    </span>
                    <span class="yl_chart_val">92%</span>
                </div>
                <!-- 其余四行结构相同，宽度依次为 86%、78%、64%、51% -->
            </div>
        </div>
    </div>
</section>

/* CSS：两栏栅格 + 条形图 */
.yl_feature { display: grid; grid-template-columns: 1.02fr 1fr; gap: 50px; align-items: center; }
.yl_feature_media img { width: 100%; aspect-ratio: 4 / 3; object-fit: cover; }
.yl_chart_row {
    display: grid; grid-template-columns: 92px 1fr 42px;
    align-items: center; gap: 12px; margin-bottom: 9px; font-size: 12px;
}
.yl_chart_track { height: 12px; background: #ece5d8; }
.yl_chart_fill { display: block; height: 100%; background: var(--yl_coral); }
.yl_chart_row:nth-child(4) .yl_chart_fill { background: var(--yl_blue); }
.yl_chart_row:nth-child(5) .yl_chart_fill { background: var(--yl_yellow); }''')

sub("网页元素3核心代码（列表）")
code_block('''<!-- HTML 结构：有序列表 -->
<ol class="yl_list">
    <li class="yl_list_item">
        <div class="yl_date"><strong>26</strong><span>10月 · 周一</span></div>
        <div class="yl_list_body">
            <h3>读书节开幕式 · 主题分享</h3>
            <p><i class="iconfont icon-shouye"></i>图书馆三楼报告厅　19:00 — 20:30</p>
        </div>
        <span class="yl_status yl_status_hot">可报名</span>
    </li>
    <!-- 其余三条结构相同 -->
</ol>

/* CSS：三列行布局 + 状态标签 */
.yl_list { list-style: none; padding: 0; border-top: 1px solid var(--yl_line); }
.yl_list_item {
    display: grid; grid-template-columns: 84px 1fr auto;
    align-items: center; gap: 26px; padding: 24px 8px;
    border-bottom: 1px solid var(--yl_line);
}
.yl_date { text-align: center; border-right: 1px solid var(--yl_line); }
.yl_date strong { font-family: var(--yl_serif); font-size: 32px; color: var(--yl_coral); }
.yl_status { padding: 5px 12px; border: 1px solid var(--yl_line); font-size: 12px; }
.yl_status_hot { border-color: var(--yl_coral); background: var(--yl_coral); color: #fffdf8; }''')

# ---------- 三、测试与发布 ----------
heading("测试与发布")
para("页面在 Chrome 与 Edge 两个浏览器中打开检查，并在 1200px 桌面视口与 390px 手机视口下查看布局：桌面端三列卡片、左右两栏与列表三列均保持完整；手机视口下卡片转为单列，左图右文改为上图下文，列表的状态标签换行到正文下方，均无横向滚动条与内容溢出。测试中发现窄屏下标题与列表右侧文字贴近边缘，已在 680px 断点中补充内边距规则。")

picture(os.path.join(SHOTS, "test_desktop_chrome.png"), 14.4, "图 4-1  Chrome 浏览器 · 1200px 桌面视口")
picture(os.path.join(SHOTS, "test_desktop_edge.png"), 14.4, "图 4-2  Edge 浏览器 · 1200px 桌面视口")
picture(os.path.join(SHOTS, "test_mobile_chrome.png"), 6.6, "图 4-3  Chrome 浏览器 · 390px 手机视口")

label("网页元素一发布图（如果有）：", 6)
para("本次项目源码与素材已整理在同一文件夹内，可直接在浏览器中打开运行；因暂未申请服务器空间，尚未发布到 gitee / github 等平台，故此处省略发布截图。")

# ---------- 四、项目总结 ----------
heading("项目总结")
para("本次项目最大的收获是把“看图还原”拆成了可执行的步骤：先量尺寸（内容区 1124px、卡片 360px、间距 22px、配图 3:2 与 4:3），再确定色板与字体，最后才写结构。顺序反过来很容易在细节上反复返工。")
para("第二个体会是布局方式的选择。三个元素全部用 Grid 完成主结构，卡片内部再用 Flex 控制“文字块底边对齐”，比过去用 float 加清除浮动简单得多；响应用媒体查询改列数即可，不需要重写结构。条形图完全用嵌套 span 的宽度绘制，让我理解了“图表不一定非要图片”。")
para("教训主要有两点：一是最初的左图右文与第二张卡片重复使用同一张配图，视觉上很单调，后来单独补了一张临海桌面的图；二是窄屏下直接裁切截图时出现过文字贴边被切的情况，说明断点里的内边距要单独确认，不能只靠桌面效果判断。")
para("AI 使用说明：本项目在 Trae 中使用 DeepSeek-V4.1-Flash 模型辅助完成。AI 参与了页面结构草拟、CSS 编写与响应式断点调整，并根据提示词生成了三张书单配图与一张共读配图；设计稿、实现图、不同浏览器与视口的测试截图均由本机 Chrome / Edge 实际渲染后截取，代码逐段人工检查并修改了图片重复、窄屏贴边等问题。模型与提示词记录见项目目录下的 zheng_prompt.md。")

doc.save(OUT_DOC)
print("saved", OUT_DOC)
